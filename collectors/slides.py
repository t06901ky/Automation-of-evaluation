"""Google Slides / PowerPoint (.pptx) からアクションアイテム目標を取得する。

ネイティブ Google Slides → Slides API で読む
アップロード済み .pptx   → Drive API でダウンロード → python-pptx でパース
"""

from __future__ import annotations

from pathlib import Path

from collectors.auth import build_service

SLIDES_SCOPES = ["https://www.googleapis.com/auth/presentations.readonly"]
DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


# ============================================================
# 自動判定
# ============================================================


def _is_pptx(sa_file: Path, file_id: str) -> bool:
    drive = build_service("drive", "v3", sa_file, DRIVE_SCOPES)
    meta = drive.files().get(fileId=file_id, fields="mimeType", supportsAllDrives=True).execute()
    return meta.get("mimeType") == PPTX_MIME


def _download_pptx(sa_file: Path, file_id: str) -> bytes:
    drive = build_service("drive", "v3", sa_file, DRIVE_SCOPES)
    return drive.files().get_media(fileId=file_id).execute()


# ============================================================
# pptx パース
# ============================================================


def _pptx_slide_text(data: bytes, target_slide_index: int | None = None) -> str:
    """pptx バイナリから指定スライド (or 全スライド) のテキストを抽出。"""
    import io

    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    all_text: list[str] = []

    for i, slide in enumerate(prs.slides):
        texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if t:
                        texts.append(t)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    texts.append("\t".join(cells))

        if not texts:
            continue

        slide_text = "\n".join(texts)
        if target_slide_index is not None and i == target_slide_index:
            return slide_text

        all_text.append(f"[Slide {i + 1}]\n{slide_text}")

    return "\n\n".join(all_text)


# ============================================================
# Slides API (ネイティブ Google Slides 用)
# ============================================================


def _slides_service(sa_file: Path):
    return build_service("slides", "v1", sa_file, SLIDES_SCOPES)


def _native_slide_text(sa_file: Path, presentation_id: str, target_slide_id: str) -> str:
    svc = _slides_service(sa_file)
    presentation = svc.presentations().get(presentationId=presentation_id).execute()

    for slide in presentation.get("slides", []):
        if slide.get("objectId") == target_slide_id:
            return _extract_text_from_slide(slide)

    all_text: list[str] = []
    for slide in presentation.get("slides", []):
        text = _extract_text_from_slide(slide)
        if text.strip():
            slide_id = slide.get("objectId", "?")
            all_text.append(f"[Slide {slide_id}]\n{text}")
    return "\n\n".join(all_text)


def _extract_text_from_slide(slide: dict) -> str:
    texts: list[str] = []
    for element in slide.get("pageElements", []):
        shape = element.get("shape")
        if not shape:
            table = element.get("table")
            if table:
                texts.append(_extract_text_from_table(table))
            continue
        text_content = shape.get("text")
        if not text_content:
            continue
        for text_element in text_content.get("textElements", []):
            text_run = text_element.get("textRun")
            if text_run:
                texts.append(text_run.get("content", ""))
    return "".join(texts)


def _extract_text_from_table(table: dict) -> str:
    rows: list[str] = []
    for table_row in table.get("tableRows", []):
        cells: list[str] = []
        for cell in table_row.get("tableCells", []):
            cell_text_parts: list[str] = []
            text_content = cell.get("text")
            if text_content:
                for te in text_content.get("textElements", []):
                    tr = te.get("textRun")
                    if tr:
                        cell_text_parts.append(tr.get("content", "").strip())
            cells.append(" ".join(cell_text_parts))
        rows.append("\t".join(cells))
    return "\n".join(rows)


# ============================================================
# 公開 API (自動判定)
# ============================================================


def fetch_slide_text(
    sa_file: Path,
    presentation_id: str,
    target_slide_id: str,
    target_slide_index: int | None = None,
) -> str:
    """スライドからテキストを取得。pptx なら自動でダウンロード + パース。"""
    if _is_pptx(sa_file, presentation_id):
        data = _download_pptx(sa_file, presentation_id)
        return _pptx_slide_text(data, target_slide_index=target_slide_index)
    return _native_slide_text(sa_file, presentation_id, target_slide_id)
