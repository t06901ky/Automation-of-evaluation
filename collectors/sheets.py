"""Google Sheets / Excel (.xlsx) からグレード定義と KPI データを取得する。

ネイティブ Google Sheets → Sheets API で読む
アップロード済み .xlsx  → Drive API でダウンロード → openpyxl でパース
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import openpyxl

from collectors.auth import build_service

SHEETS_SCOPES = ["https://www.googleapis.com/auth/spreadsheets.readonly"]
DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


# ============================================================
# 自動判定: ネイティブ Sheets か xlsx か
# ============================================================


def _is_xlsx(sa_file: Path, file_id: str) -> bool:
    drive = build_service("drive", "v3", sa_file, DRIVE_SCOPES)
    meta = drive.files().get(fileId=file_id, fields="mimeType", supportsAllDrives=True).execute()
    return meta.get("mimeType") == XLSX_MIME


def _download_xlsx(sa_file: Path, file_id: str) -> bytes:
    drive = build_service("drive", "v3", sa_file, DRIVE_SCOPES)
    return drive.files().get_media(fileId=file_id).execute()


def _xlsx_to_rows(data: bytes, sheet_name: str | None = None, gid: int | None = None) -> list[list[Any]]:
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = None
    if gid is not None:
        # gid でシートを探す (openpyxl は gid を直接持たないので index で代替)
        # gid は通常シートの並び順に対応しないが、ここでは index として試す
        for i, name in enumerate(wb.sheetnames):
            if i == gid or str(gid) in name:
                ws = wb[name]
                break
    if ws is None and sheet_name:
        for name in wb.sheetnames:
            if sheet_name.lower() in name.lower():
                ws = wb[name]
                break
    if ws is None:
        ws = wb.active or wb[wb.sheetnames[0]]

    rows: list[list[Any]] = []
    for row in ws.iter_rows(values_only=True):
        # 全 None の行はスキップ
        if all(c is None for c in row):
            continue
        rows.append([c if c is not None else "" for c in row])
    wb.close()
    return rows


# ============================================================
# Sheets API (ネイティブ Google Sheets 用)
# ============================================================


def _sheets_service(sa_file: Path):
    return build_service("sheets", "v4", sa_file, SHEETS_SCOPES)


def _fetch_range(sa_file: Path, spreadsheet_id: str, range_a1: str) -> list[list[Any]]:
    svc = _sheets_service(sa_file)
    resp = (
        svc.spreadsheets()
        .values()
        .get(spreadsheetId=spreadsheet_id, range=range_a1)
        .execute()
    )
    return resp.get("values", [])


# ============================================================
# グレード定義 (ネイティブ Sheets 想定)
# ============================================================


def fetch_grade_definition(
    sa_file: Path,
    spreadsheet_id: str,
    sheet_name: str = "グレード",
) -> str:
    """グレード定義シートをテキストとして返す。"""
    try:
        rows = _fetch_range(sa_file, spreadsheet_id, f"'{sheet_name}'!A1:Z500")
    except Exception:
        rows = []
    if not rows:
        try:
            rows = _fetch_range(sa_file, spreadsheet_id, "A1:Z500")
        except Exception:
            rows = []
    lines = ["\t".join(str(c) for c in row) for row in rows]
    return "\n".join(lines)


def extract_grade_section(full_text: str, grade: str) -> str:
    """特定グレードの行だけを抽出。見つからなければ全体を返す。"""
    if not full_text:
        return ""
    lines = full_text.splitlines()

    # ヘッダ行 (最初の数行) を保持
    header_lines = []
    for line in lines:
        if line and line[0].isdigit() and "\t" in line[:3]:
            break
        header_lines.append(line)

    # 対象グレードの行を「グレード番号\t」で始まる行から次のグレードまで
    grade_lines = []
    capturing = False
    for line in lines:
        if not capturing and line.startswith(grade + "\t"):
            capturing = True
            grade_lines.append(line)
            continue
        if capturing:
            if len(line) > 0 and line[0].isdigit() and "\t" in line[:3]:
                break
            grade_lines.append(line)

    if grade_lines:
        return "\n".join(header_lines + grade_lines)
    return full_text


# ============================================================
# KPI (xlsx 対応)
# ============================================================


def fetch_kpi_data(
    sa_file: Path,
    spreadsheet_id: str,
    sheet_name: str | None = None,
    gid: int | None = None,
) -> list[dict[str, Any]]:
    """KPI データを dict リストで返す。xlsx ならダウンロードしてパース。"""
    rows = _fetch_rows_auto(sa_file, spreadsheet_id, sheet_name, gid)
    return _rows_to_dicts(rows)


def fetch_kpi_raw_text(
    sa_file: Path,
    spreadsheet_id: str,
    sheet_name: str | None = None,
    gid: int | None = None,
) -> str:
    """KPI シートをテキストとして返す (LLM 用)。"""
    rows = _fetch_rows_auto(sa_file, spreadsheet_id, sheet_name, gid)
    lines = ["\t".join(str(c) for c in row) for row in rows]
    return "\n".join(lines)


def _fetch_rows_auto(
    sa_file: Path,
    spreadsheet_id: str,
    sheet_name: str | None,
    gid: int | None,
) -> list[list[Any]]:
    """ネイティブ Sheets か xlsx かを自動判定して行を返す。"""
    if _is_xlsx(sa_file, spreadsheet_id):
        data = _download_xlsx(sa_file, spreadsheet_id)
        return _xlsx_to_rows(data, sheet_name=sheet_name, gid=gid)
    # ネイティブ Sheets
    range_name = f"'{sheet_name}'!A1:Z500" if sheet_name else "A1:Z500"
    try:
        return _fetch_range(sa_file, spreadsheet_id, range_name)
    except Exception:
        return _fetch_range(sa_file, spreadsheet_id, "A1:Z500")


# ============================================================
# ユーティリティ
# ============================================================


def _rows_to_dicts(rows: list[list[Any]]) -> list[dict[str, Any]]:
    if not rows:
        return []
    header = [str(h).strip() for h in rows[0]]
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        padded = list(row) + [""] * (len(header) - len(row))
        out.append(dict(zip(header, padded, strict=False)))
    return out
